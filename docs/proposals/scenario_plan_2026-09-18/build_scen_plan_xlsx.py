# -*- coding: utf-8 -*-
"""Build the 14-scenario quantification plan workbook from scen_plan_content.py (model side)
and scen_plan_research.py (scenario rows, assumption rows, sources, open questions).
Usage: python build_scen_plan_xlsx.py [out.xlsx]"""
import sys, os, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from scen_plan_content import BASELINE, DC_BY_REGION, EXISTING, MODEL_FACTS
import importlib
RES = importlib.import_module(os.environ.get('SCEN_RESEARCH', 'scen_plan_research'))
SCENARIOS, ASSUMPTIONS, SOURCES, OPEN_QUESTIONS = RES.SCENARIOS, RES.ASSUMPTIONS, RES.SOURCES, RES.OPEN_QUESTIONS
V2 = 'v2' in os.environ.get('SCEN_RESEARCH', '')

OUT = sys.argv[1] if len(sys.argv) > 1 else 'NA_scenario_plan_2026-09.xlsx'
HDR = PatternFill('solid', fgColor='1F3864'); HFONT = Font(bold=True, color='FFFFFF')
FLAG = {'GREEN': 'C6EFCE', 'YELLOW': 'FFEB9C', 'RED': 'FFC7CE'}
WRAP = Alignment(wrap_text=True, vertical='top')
thin = Side(style='thin', color='BFBFBF'); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def sheet(wb, name, header, rows, widths, flag_col=None, freeze='A2'):
    ws = wb.create_sheet(name)
    ws.append(header)
    for c in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HDR; cell.font = HFONT; cell.alignment = WRAP; cell.border = BORDER
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(header)):
        for cell in row:
            cell.alignment = WRAP; cell.border = BORDER
        if flag_col is not None:
            v = str(row[flag_col].value or '').upper()
            for k, col in FLAG.items():
                if v.startswith(k):
                    row[flag_col].fill = PatternFill('solid', fgColor=col)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    return ws


wb = Workbook(); wb.remove(wb.active)

# --- README
ws = wb.create_sheet('README')
lines = [
    ('NA model scenario quantification plan' + (' - v2' if V2 else ''), True),
    (f'Prepared {datetime.date.today().isoformat()} for the 14 requested SE North America scenarios / sensitivities. '
     'Status: PLAN (nothing implemented). Numbers come from web-sourced reports (2025-2026 vintage) plus the model baseline; '
     'every number in "Assumptions" carries a source id or is explicitly marked as an assumption.', False),
    ('', False),
    ('Sheets', True),
    ('Overview - one row per scenario: direction, model levers, implementation route, REPRESENTABILITY flag (can the model do it), EVIDENCE strength (how well sourced - informational, not a value judgement), overlap with existing sensitivities, proposed variants incl. a stretch what-if, effort.', False),
    ('Assumptions - the quantified inputs per scenario (parameter, region, years, value, unit, derivation, source id, confidence, invented flag). This is what gets implemented.', False),
    ('Feasibility - concerns per scenario with a Type column: model = representation limitation (drives the flag), data = evidence note (does not drive the flag). Flags: GREEN = existing levers; YELLOW = needs a new lever/technology or an approximation; RED = cannot be represented. v2 (2026-09-18): judgement calls and far-fetched what-ifs are legitimate scenario content and are not down-rated; every scenario has a stretch variant; fusion capex = user input.', False),
    ('Sources - reference list (id, organisation, title, date, URL, used for).', False),
    ('Model baseline - what the model currently assumes (FEL26 base demand, funnels, costs, policy), so the deltas are visible.', False),
    ('Existing sensitivities - the v6 sensitivity set the new scenarios can reuse.', False),
    ('Open questions - decisions needed from the requesters before implementation.', False),
    ('', False),
    ('Conventions', True),
    ('Regions: California, WECC (rest of US West), ERCOT, SPP, MISO, SERC (US Southeast), PJM, NewYork, NewEngland, Canada (one node). '
     'Years 2025-2040 annual. Demand in TWh busbar. Costs in EUR2023 (1 USD = 0.925 EUR). Gas prices in USD/MMBtu where quoted; 1 USD/MMBtu = 0.877 EUR/GJ.', False),
    ('Effort: S = data rows via existing scripts (<1 day incl. run); M = new script/lever or new technology (1-3 days); L = new model feature or FEL demand-model rerun by the FEL team (>3 days / external dependency).', False),
    ('Tags from the request: [Tailwind] upside to load / SE market, [Headwind] downside, [Mixed].', False),
]
for t, b in lines:
    ws.append([t])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=b, size=12 if b else 11)
    ws.cell(row=ws.max_row, column=1).alignment = WRAP
ws.column_dimensions['A'].width = 140

# --- Overview
sheet(wb, 'Overview',
      ['ID', 'Scenario', 'Tag', 'What it means for the model', 'Model levers (parameters)', 'Implementation route',
       'Representability', 'Representability summary', 'Evidence strength (info)', 'Overlap with existing sensitivities', 'Proposed variants (incl. stretch)', 'Effort', 'Key external anchors'],
      [(s['id'], s['name'], s['tag'], s['meaning'], s['levers'], s['route'], s['flag'], s['flag_summary'], s.get('evidence', ''),
        s['overlap'], s['variants'], s['effort'], s['anchors']) for s in SCENARIOS],
      [5, 26, 11, 45, 45, 38, 12, 45, 28, 32, 48, 8, 45], flag_col=6, freeze='C2')

# --- Assumptions
sheet(wb, 'Assumptions',
      ['ID', 'Scenario', 'Variant', 'Parameter / lever', 'Region(s)', 'Years', 'Value', 'Unit',
       'Basis / derivation', 'Source id(s)', 'Confidence', 'Invented?'],
      [tuple(a) for a in ASSUMPTIONS],
      [5, 22, 14, 34, 22, 14, 22, 14, 60, 16, 11, 10], flag_col=None, freeze='D2')
ws = wb['Assumptions']
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    v = str(row[11].value or '').lower()
    if v.startswith('yes'):
        row[11].fill = PatternFill('solid', fgColor=FLAG['RED'])
    elif v.startswith('partly'):
        row[11].fill = PatternFill('solid', fgColor=FLAG['YELLOW'])
    elif v.startswith('user'):
        row[11].fill = PatternFill('solid', fgColor='BDD7EE')
    c = str(row[10].value or '').lower()
    if c.startswith('high'):
        row[10].fill = PatternFill('solid', fgColor=FLAG['GREEN'])
    elif c.startswith('low'):
        row[10].fill = PatternFill('solid', fgColor=FLAG['RED'])
    elif c.startswith('med'):
        row[10].fill = PatternFill('solid', fgColor=FLAG['YELLOW'])

# --- Feasibility
feas = []
for s in SCENARIOS:
    for c in s['concerns']:
        if len(c) == 4:
            sev, typ, why, fix = c
        else:
            sev, why, fix = c; typ = ''
        feas.append((s['id'], s['name'], sev, typ, why, fix))
sheet(wb, 'Feasibility', ['ID', 'Scenario', 'Severity', 'Type', 'Concern', 'Workaround / proposal'], feas, [5, 26, 10, 8, 70, 70], flag_col=2)

# --- Sources
sheet(wb, 'Sources', ['Source id', 'Organisation', 'Title', 'Date', 'URL', 'Used for'], SOURCES, [10, 26, 60, 12, 60, 30])

# --- Model baseline
rows = [('FEL26 base demand as loaded (busbar TWh)', '', '', '', '', '')]
rows.append(('Category', 'Scope', '2025', '2030', '2035', '2040'))
for f, d in BASELINE.items():
    for sc, v in d.items():
        rows.append((f, 'US 9 pools' if sc == 'US' else 'Canada', v[2025], v[2030], v[2035], v[2040]))
rows.append(('', '', '', '', '', ''))
rows.append(('Power_DataCenter by region (TWh)', 'Region', '2025', '2030', '2040', ''))
for r, v in DC_BY_REGION.items():
    rows.append(('', r, v[0], v[1], v[2], ''))
rows.append(('', '', '', '', '', ''))
rows.append(('Model facts', '', '', '', '', ''))
for k, v in MODEL_FACTS:
    rows.append((k, v, '', '', '', ''))
sheet(wb, 'Model baseline', ['Item', 'Detail', 'a', 'b', 'c', 'd'], rows, [34, 110, 10, 10, 10, 10], freeze='A2')

# --- Existing sensitivities
sheet(wb, 'Existing sensitivities', ['Sensitivity (v6 set)', 'Definition'], list(EXISTING.items()), [34, 110])

# --- Open questions
sheet(wb, 'Open questions', ['#', 'Scenario', 'Question', 'Why it matters'], OPEN_QUESTIONS, [4, 22, 70, 60])

wb.save(OUT)
print('saved', OUT)
